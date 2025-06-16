from docx import Document


def fill_template(template_path, output_path, placeholders):
    """
    Fills placeholders in a Word template and saves the result.
    :param template_path: Path to the Word template (.docx).
    :param output_path: Path to save the completed document (.docx).
    :param placeholders: Dictionary with placeholder-text values (e.g., {NAME}: John).
    """
    doc = Document(template_path)  # Load the Word template

    # Replace placeholders in all paragraphs
    for paragraph in doc.paragraphs:
        for placeholder, value in placeholders.items():
            if placeholder in paragraph.text:
                paragraph.text = paragraph.text.replace(placeholder, value)

    # Save the modified document
    doc.save(output_path)